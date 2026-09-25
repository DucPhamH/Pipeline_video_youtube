"""Xuất Excel + review-all."""
from crawl.application.excel_export import content_disposition
from crawl.application.use_cases import RawTextStorage
from crawl.domain.entities import Chapter, ChapterStatus, Novel, NovelLifecycle
from crawl.infrastructure.persistence.repositories import (
    SqlAlchemyChapterRepository,
    SqlAlchemyNovelRepository,
)
from platform_.config import config
from platform_.db import SessionLocal


def _seed_novel_with_chapters(*, reviewed: bool, cleaned: bool, suffix: str):
    db = SessionLocal()
    storage = RawTextStorage(config.raw_dir)
    try:
        novel_repo = SqlAlchemyNovelRepository(db)
        chapter_repo = SqlAlchemyChapterRepository(db)
        novel = novel_repo.add(
            Novel(
                id=None,
                title=f"Export Novel {suffix}",
                source_key="demo_local",
                source_url=f"http://demo.local/export-{suffix}",
                lifecycle_status=NovelLifecycle.FULLY_CRAWLED,
                last_chapter_index=2,
                total_chapters=2,
            )
        )
        ids = []
        for i in (1, 2):
            path = storage.save(novel.id, i, f"raw chapter {i} content line\n")
            if cleaned:
                storage.save_cleaned(path, f"cleaned chapter {i}\n")
            ch = chapter_repo.add(
                Chapter(
                    id=None,
                    novel_id=novel.id,
                    chapter_index=i,
                    title=f"Ch {i}",
                    source_url=f"http://demo.local/{suffix}/c{i}",
                    raw_path=path,
                    status=ChapterStatus.CRAWLED,
                    reviewed=reviewed,
                )
            )
            ids.append(ch.id)
        return novel.id, ids, storage
    finally:
        db.close()


def test_export_workbook_and_status(client):
    nid, _, _ = _seed_novel_with_chapters(reviewed=False, cleaned=True, suffix="wb2")
    st = client.get(f"/api/crawl/novels/{nid}/export-status")
    assert st.status_code == 200
    body = st.json()
    assert body["crawled"] == 2
    assert body["cleaned"] == 2
    assert body["can_export_workbook"] is True
    assert body["can_export_txt"] is True
    assert body["can_export_epub"] is True
    assert body["can_export_bundle"] is True
    assert "can_export_per_chapter" not in body

    r = client.get(f"/api/crawl/novels/{nid}/export.xlsx")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"

    txt = client.get(f"/api/crawl/novels/{nid}/export.txt")
    assert txt.status_code == 200
    assert b"cleaned chapter 1" in txt.content

    epub = client.get(f"/api/crawl/novels/{nid}/export.epub")
    assert epub.status_code == 200
    assert epub.content[:2] == b"PK"

    z = client.get(f"/api/crawl/novels/{nid}/export.zip")
    assert z.status_code == 200
    assert z.content[:2] == b"PK"

    batch = client.post(
        "/api/crawl/novels/export-batch",
        json={"novel_ids": [nid], "format": "txt"},
    )
    assert batch.status_code == 200
    assert batch.content[:2] == b"PK"

    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Tiêu đề", "Nội dung")
    assert rows[1][0] == "Ch 1"
    assert str(rows[1][1]).startswith("Ch 1")
    assert "cleaned chapter 1" in str(rows[1][1])


def test_export_filename_unicode_header():
    cd = content_disposition("太荒吞天诀.xlsx", fallback="novel-1.xlsx")
    assert "novel-1.xlsx" in cd
    assert "filename*=UTF-8''" in cd
    cd.encode("latin-1")


def test_review_all_and_skips_without_cleaned(client):
    nid, _, _ = _seed_novel_with_chapters(reviewed=False, cleaned=True, suffix="ra2")
    rev = client.post(f"/api/crawl/novels/{nid}/review-all")
    assert rev.status_code == 200
    assert rev.json()["chapters_reviewed"] == 2

    nid2, _, _ = _seed_novel_with_chapters(reviewed=False, cleaned=False, suffix="ra3")
    rev2 = client.post(f"/api/crawl/novels/{nid2}/review-all")
    assert rev2.status_code == 200
    assert rev2.json()["chapters_reviewed"] == 0
